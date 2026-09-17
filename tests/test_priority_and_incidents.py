import pytest
import io
import openpyxl
from src.coordination.priority import PriorityClass, PriorityEvaluator, PriorityWeights, PriorityEvaluation
from src.coordination.conflict_resolution import NegotiationProtocol, WaitForGraph
from src.coordination.incidents import IncidentManager, IncidentSeverity, IncidentStatus
from src.robots.amr import AMRRobot
from src.simulation.simulator import DecentralizedFleetSimulator, SimulationConfig
from src.tasks.task import Task
from src.warehouse.warehouse import Warehouse
from src.reporting.excel_export import export_simulation_to_excel


def test_priority_evaluation_deterministic():
    evaluator = PriorityEvaluator()
    eval1 = evaluator.evaluate(
        robot_id="AMR-001",
        carrying_package=True,
        task_priority=3,
        battery=85.0,
        time_step=10,
        deadline=30,
        remaining_distance=8,
    )
    eval2 = evaluator.evaluate(
        robot_id="AMR-001",
        carrying_package=True,
        task_priority=3,
        battery=85.0,
        time_step=10,
        deadline=30,
        remaining_distance=8,
    )
    assert eval1.score == eval2.score
    assert eval1.priority_class == eval2.priority_class
    assert eval1.factors == eval2.factors
    assert eval1.score > 0
    assert "CarryingCargo" in eval1.explanation


def test_priority_classes_ordering():
    evaluator = PriorityEvaluator()
    low = evaluator.evaluate("AMR-001", False, 1, 80.0)
    normal = evaluator.evaluate("AMR-001", False, 2, 80.0)
    high = evaluator.evaluate("AMR-001", False, 3, 80.0)
    critical = evaluator.evaluate("AMR-001", False, 4, 80.0)
    emergency = evaluator.evaluate("AMR-001", False, 5, 80.0, is_emergency=True)

    assert low.score < normal.score < high.score < critical.score < emergency.score
    assert emergency.priority_class == PriorityClass.EMERGENCY
    assert critical.priority_class == PriorityClass.CRITICAL


def test_loaded_vs_empty_narrow_aisle_arbitration():
    evaluator = PriorityEvaluator()
    # Loaded AMR with normal task
    eval_loaded = evaluator.evaluate("AMR-001", carrying_package=True, task_priority=2, battery=80.0, remaining_distance=4)
    # Empty AMR with same priority task
    eval_empty = evaluator.evaluate("AMR-002", carrying_package=False, task_priority=2, battery=90.0, remaining_distance=4)

    winner, loser, reason = evaluator.arbitrate_contention(eval_loaded, eval_empty, "AMR-001", "AMR-002")
    assert winner == "AMR-001"
    assert loser == "AMR-002"
    assert "cargo" in reason.lower() or "right-of-way" in reason.lower()


def test_sla_deadline_priority_escalation():
    evaluator = PriorityEvaluator()
    # Task A has plenty of slack (deadline 100, time 10, dist 10)
    eval_relaxed = evaluator.evaluate("AMR-001", False, 2, 80.0, time_step=10, deadline=100, remaining_distance=10)
    # Task B is at risk (deadline 15, time 10, dist 10)
    eval_urgent = evaluator.evaluate("AMR-002", False, 2, 80.0, time_step=10, deadline=15, remaining_distance=10)

    assert eval_urgent.score > eval_relaxed.score
    assert eval_urgent.factors["sla_risk"] > eval_relaxed.factors["sla_risk"]


def test_emergency_task_preemption_in_simulator():
    warehouse = Warehouse(width=12, height=12)
    sim = DecentralizedFleetSimulator(warehouse, SimulationConfig(seed=42, robot_count=2, task_count=0))
    sim.initialize(create_tasks=False)

    # Robot 1 is busy with a low priority task (not yet carrying package)
    r1 = sim.robots["AMR-001"]
    r1.position = (2, 2)
    low_task = Task(task_id="T-LOW", pickup=(4, 4), destination=(8, 8), priority=1)
    sim.add_task(low_task)
    low_task.assigned_robot = r1.robot_id
    low_task.status = "in_progress"
    r1.current_task = low_task.task_id

    # Robot 2 is busy with low priority task
    r2 = sim.robots["AMR-002"]
    r2.position = (9, 9)
    low_task2 = Task(task_id="T-LOW2", pickup=(9, 8), destination=(1, 1), priority=2)
    sim.add_task(low_task2)
    low_task2.assigned_robot = r2.robot_id
    low_task2.status = "in_progress"
    r2.current_task = low_task2.task_id

    # Emergency order arrives at (3, 3)
    emergency = Task(task_id="T-EMERGENCY", pickup=(3, 3), destination=(10, 10), priority=5)
    emergency.metadata["is_emergency"] = True
    sim.add_task(emergency)

    selection = sim._select_robot_for_task(emergency)
    assert selection is not None
    robot_id, reason = selection
    assert "PREEMPTION" in reason
    # Verify low priority task was unassigned and re-queued
    assert low_task.status == "pending" or low_task2.status == "pending"
    # Verify decision logger recorded the preemption
    recent_decisions = sim.decision_logger.get_recent(5)
    assert any(d["category"] == "TASK_PREEMPTION" for d in recent_decisions)


def test_priority_wfg_deadlock_breaking():
    warehouse = Warehouse(width=10, height=10)
    sim = DecentralizedFleetSimulator(warehouse, SimulationConfig(seed=42, robot_count=3, task_count=0))
    sim.initialize(create_tasks=False)

    robots = list(sim.robots.values())
    r1, r2, r3 = robots[0], robots[1], robots[2]

    # r1 carrying high-value package
    r1.carrying_package_id = "PKG-VIP"
    t1 = Task(task_id="T1", pickup=(1, 1), destination=(5, 5), priority=4)
    sim.add_task(t1)
    r1.current_task = t1.task_id

    # r2 carrying normal package
    r2.carrying_package_id = "PKG-NORM"
    t2 = Task(task_id="T2", pickup=(1, 1), destination=(5, 5), priority=2)
    sim.add_task(t2)
    r2.current_task = t2.task_id

    # r3 empty
    r3.carrying_package_id = None
    t3 = Task(task_id="T3", pickup=(1, 1), destination=(5, 5), priority=1)
    sim.add_task(t3)
    r3.current_task = t3.task_id

    # Create cycle: r1 -> r2 -> r3 -> r1
    sim.wait_for_graph.add_wait(r1.robot_id, r2.robot_id)
    sim.wait_for_graph.add_wait(r2.robot_id, r3.robot_id)
    sim.wait_for_graph.add_wait(r3.robot_id, r1.robot_id)

    r1.state = "WAITING"
    r2.state = "WAITING"
    r3.state = "WAITING"

    sim._detect_and_resolve_wfg_deadlocks()

    # The robot with lowest operational score (r3) must be the one commanded to reroute!
    assert r3.state == "REROUTING"
    # r1 (highest score) must NOT be rerouting
    assert r1.state == "WAITING"


def test_incident_lifecycle_and_escalation():
    mgr = IncidentManager()
    inc = mgr.raise_incident(
        incident_type="AMR_FAILURE",
        severity=IncidentSeverity.CRITICAL,
        timestamp=50,
        affected_entities={"robots": ["AMR-003"]},
        detected_by="MotorEncoder",
        decision="Attempting peer rescue",
    )
    assert inc.status == IncidentStatus.OPEN
    assert inc.severity == IncidentSeverity.CRITICAL

    # Escalate if rescue impossible
    mgr.escalate(
        incident_id=inc.id,
        why="Actuator motor physically stalled on ramp",
        blocked="Aisle 4 completely impassable",
        attempted_options="Tried reverse 2 ticks and pivot turn; 0 torque",
        human_action_required="Maintenance crew manual cart extraction",
    )
    assert inc.status == IncidentStatus.ESCALATED
    assert inc.escalation_details is not None
    assert "Maintenance crew" in inc.escalation_details["human_action_required"]

    # Resolve
    mgr.resolve(inc.id, timestamp=65, outcome_action="Cart extracted and path cleared")
    assert inc.status == IncidentStatus.RESOLVED
    assert inc.resolution_time == 65


def test_operator_override_actions():
    warehouse = Warehouse(width=10, height=10)
    sim = DecentralizedFleetSimulator(warehouse, SimulationConfig(seed=42, robot_count=2, task_count=0))
    sim.initialize(create_tasks=False)

    r1 = sim.robots["AMR-001"]
    # Pause robot
    res = sim.operator_pause_robot(r1.robot_id, reason="Maintenance inspection")
    assert res is True
    assert r1.state == "PAUSED"
    assert len(sim.operator_actions) == 1
    assert sim.operator_actions[0]["action"] == "pause_robot"

    # Resume robot
    sim.operator_resume_robot(r1.robot_id)
    assert r1.state == "IDLE"

    # Close aisle
    cell = (5, 5)
    closed = sim.operator_close_aisle(cell, reason="Oil spill cordon")
    assert closed is True
    assert not sim.warehouse.is_walkable(cell)

    # Open aisle
    opened = sim.operator_open_aisle(cell, reason="Spill cleaned")
    assert opened is True
    assert sim.warehouse.is_walkable(cell)


def test_extended_excel_export_9_sheets():
    warehouse = Warehouse(width=10, height=10)
    warehouse.add_charging_station((1, 1))
    sim = DecentralizedFleetSimulator(warehouse, SimulationConfig(seed=42, robot_count=2, task_count=2))
    sim.initialize()

    # Add an incident and operator action
    sim.incident_manager.raise_incident(
        incident_type="TEST_INCIDENT",
        severity=IncidentSeverity.MEDIUM,
        timestamp=sim.time_step,
        affected_entities={"robots": ["AMR-001"]},
    )
    sim.operator_pause_robot("AMR-001", "Test operator action")

    excel_bytes = export_simulation_to_excel(sim)
    assert len(excel_bytes) > 5000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert len(sheet_names) == 9
    assert "Fleet KPIs & Overview" in sheet_names
    assert "Tasks" in sheet_names
    assert "AMRs" in sheet_names
    assert "Decision Records" in sheet_names
    assert "Deadlocks & WFG" in sheet_names
    assert "Benchmark Comparison" in sheet_names
    assert "Priority Decisions" in sheet_names
    assert "Incidents" in sheet_names
    assert "Operator Actions" in sheet_names

    # Check rows in Priority Decisions sheet
    ws_pri = wb["Priority Decisions"]
    assert ws_pri.max_row >= 2  # header + task rows
